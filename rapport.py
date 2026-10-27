"""
Production des résultats :
- un compte rendu lisible (Markdown, s'affiche proprement sur GitHub)
- un fichier Excel mis en forme pour l'équipe
"""

from datetime import date

import pandas as pd
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import config

COLONNES_EXCEL = ["Titre", "Organisation", "Date de publication", "Date de clôture",
                  "Jours restants", "Score", "Axe", "Projet à présenter",
                  "Mots-clés trouvés", "URL"]
LARGEURS = {"Titre": 60, "Organisation": 28, "Date de publication": 14,
            "Date de clôture": 14, "Jours restants": 11, "Score": 9, "Axe": 22,
            "Projet à présenter": 26, "Mots-clés trouvés": 40, "URL": 9,
            "Descriptif": 80, "Commentaire": 40, "Pertinent ?": 16}
ENTETE = PatternFill("solid", fgColor="1F4E3D")
ROUGE = PatternFill("solid", fgColor="F8CBAD")


# ---------------------------------------------------------------
# Sélections communes
# ---------------------------------------------------------------
def offres_pertinentes(df):
    """Appels ouverts au-dessus du seuil, triés par score."""
    ouverts = df[df["Statut"] == "Ouvert"]
    return ouverts[ouverts["Score"] >= config.SEUIL_PERTINENCE].sort_values("Score", ascending=False)


def financeurs(df):
    """Organisations ayant publié des appels pertinents (ouverts ou clos)."""
    pert = df[(df["Score"] >= config.SEUIL_PERTINENCE) & (df["Organisation"] != "")]
    if pert.empty:
        return pd.DataFrame(columns=["Organisation", "Appels pertinents", "Score moyen", "Dernier appel"])
    return (
        pert.groupby("Organisation")
        .agg(**{"Appels pertinents": ("Titre", "size"),
                "Score moyen": ("Score", "mean"),
                "Dernier appel": ("Date de publication", "max")})
        .round({"Score moyen": 2})
        .sort_values(["Appels pertinents", "Score moyen"], ascending=False)
        .reset_index()
    )


# ---------------------------------------------------------------
# Compte rendu Markdown
# ---------------------------------------------------------------
def _cellule(texte):
    return str(texte).replace("|", "/").replace("\n", " ")


def _date_fr(d):
    return d.strftime("%d/%m/%Y") if pd.notna(d) else "?"


def _tableau(offres):
    lignes = ["| # | Appel | Organisation | Clôture | Jours | Score | Axe |",
              "|---|---|---|---|---|---|---|"]
    for i, (_, o) in enumerate(offres.iterrows(), start=1):
        orga = o["Organisation"] or "Non renseigné"
        lignes.append(
            f"| {i} | [{_cellule(o['Titre'])}]({o['URL']}) | {_cellule(orga)} | "
            f"{_date_fr(o['Date de clôture'])} | {int(o['Jours restants'])} | "
            f"{o['Score']:.2f} | {o['Axe']} |"
        )
    return "\n".join(lignes)


def _fiches(offres):
    blocs = []
    for _, o in offres.iterrows():
        resume = o["Descriptif"][:280].rsplit(" ", 1)[0] + "…" if len(o["Descriptif"]) > 280 else o["Descriptif"]
        pourquoi = o["Mots-clés trouvés"] or "proximité de sens uniquement"
        blocs.append(
            f"#### [{o['Titre']}]({o['URL']})\n"
            f"**{o['Organisation'] or 'Organisation non renseignée'}** · "
            f"clôture le {_date_fr(o['Date de clôture'])} ({int(o['Jours restants'])} jours) · "
            f"score {o['Score']:.2f}\n\n"
            f"*Pourquoi* · axe {o['Axe']}, mots-clés {pourquoi}\n\n"
            f"> {resume}\n"
        )
    return "\n".join(blocs)


def creer_compte_rendu(df, chemin):
    pert = offres_pertinentes(df).head(config.NB_OFFRES_RAPPORT)
    urgents = pert[pert["Jours restants"] <= config.URGENCE_JOURS].sort_values("Jours restants")
    autres = pert.drop(urgents.index)

    md = [
        f"# Veille appels à projets · E2C Paris",
        f"*Compte rendu du {date.today():%d/%m/%Y} · source Carenews*",
        "",
        "## En bref",
        "",
        "| Indicateur | Valeur |",
        "|---|---|",
        f"| Offres analysées | {len(df)} |",
        f"| Appels encore ouverts | {(df['Statut'] == 'Ouvert').sum()} |",
        f"| Appels ouverts pertinents (score ≥ {config.SEUIL_PERTINENCE}) | {len(offres_pertinentes(df))} |",
        f"| Dont clôture sous {config.URGENCE_JOURS} jours | {len(urgents)} |",
        "",
    ]

    md += [f"## À traiter en priorité (clôture sous {config.URGENCE_JOURS} jours)", ""]
    md += [_tableau(urgents) if not urgents.empty else "*Aucun appel pertinent n'arrive à échéance.*", ""]

    for projet in pert["Projet à présenter"].dropna().unique():
        sel = autres[autres["Projet à présenter"] == projet]
        if sel.empty:
            continue
        md += [f"## Pour le projet « {projet} »", "", _tableau(sel), ""]

    if not pert.empty:
        md += ["## Détail des offres", "", _fiches(pert), ""]

    fin = financeurs(df).head(10)
    if not fin.empty:
        md += ["## Financeurs les plus actifs sur les thèmes de l'E2C", "",
               "| Organisation | Appels pertinents | Score moyen |", "|---|---|---|"]
        md += [f"| {_cellule(r['Organisation'])} | {r['Appels pertinents']} | {r['Score moyen']:.2f} |"
               for _, r in fin.iterrows()]
        md += [""]

    md += [
        "## Méthode",
        "",
        "Chaque offre reçoit un score de 0 à 1 qui combine deux mesures à parts égales. "
        "Le score mots-clés repère les termes liés aux missions de l'E2C (insertion, jeunes de 16 à 25 ans, "
        "égalité des chances, partenaires). Le score BERT mesure la proximité de sens entre l'offre et "
        "les trois axes de l'E2C (insertion des jeunes, engagement citoyen, égalité et inclusion). "
        "Les réglages se trouvent dans `config.py`.",
        "",
    ]

    with open(chemin, "w", encoding="utf-8") as f:
        f.write("\n".join(md))


# ---------------------------------------------------------------
# Excel
# ---------------------------------------------------------------
def _mise_en_forme(ws):
    entetes = {c.value: c.column_letter for c in ws[1]}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for nom, lettre in entetes.items():
        ws.column_dimensions[lettre].width = LARGEURS.get(nom, 18)

    fin = ws.max_row
    if fin < 2:
        return
    for nom in ("Date de publication", "Date de clôture", "Dernier appel"):
        if nom in entetes:
            for c in ws[entetes[nom]][1:]:
                c.number_format = "DD/MM/YYYY"
    if "URL" in entetes:
        for c in ws[entetes["URL"]][1:]:
            if c.value:
                c.hyperlink = c.value
                c.value = "Voir"
                c.style = "Hyperlink"
    if "Score" in entetes:
        l = entetes["Score"]
        ws.conditional_formatting.add(f"{l}2:{l}{fin}", ColorScaleRule(
            start_type="num", start_value=0, start_color="FFFFFF",
            end_type="num", end_value=1, end_color="63BE7B"))
    if "Jours restants" in entetes:
        l = entetes["Jours restants"]
        ws.conditional_formatting.add(f"{l}2:{l}{fin}", CellIsRule(
            operator="between", formula=["0", str(config.URGENCE_JOURS)], fill=ROUGE))


def creer_excel(df, chemin):
    a_traiter = offres_pertinentes(df).sort_values(["Jours restants", "Score"], ascending=[True, False])
    toutes = df.sort_values("Score", ascending=False)
    avis = a_traiter[["Titre", "Organisation", "URL"]].assign(**{"Pertinent ?": "", "Commentaire": ""})

    with pd.ExcelWriter(chemin, engine="openpyxl") as w:
        a_traiter[COLONNES_EXCEL].to_excel(w, sheet_name="À traiter", index=False)
        toutes[COLONNES_EXCEL + ["Statut", "Descriptif"]].to_excel(w, sheet_name="Toutes les offres", index=False)
        financeurs(df).to_excel(w, sheet_name="Financeurs", index=False)
        avis.to_excel(w, sheet_name="Avis", index=False)

        for ws in w.book.worksheets:
            _mise_en_forme(ws)

        ws_avis = w.book["Avis"]
        dv = DataValidation(type="list", formula1='"Oui,Non,Déjà candidaté"', allow_blank=True)
        ws_avis.add_data_validation(dv)
        dv.add(f"D2:D{max(ws_avis.max_row, 2)}")
